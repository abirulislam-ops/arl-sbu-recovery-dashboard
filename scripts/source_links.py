#!/usr/bin/env python3
"""Collect every data link the SBU tabs carry, and check whether it is readable.

Source: the Data Request sheet (one tab per SBU). Column F now holds either an
ERP table reference or a pasted Google Drive / Sheets link. This script never
copies numbers — it only records what is linked and whether we can open it, so
the dashboard can show honest coverage per SBU.

Output: data/source_links.json
"""
import io, json, os, re, sys, urllib.request

SHEET = "1JVDgd6KG2K0-EUJu8a8fZoK0yLDa0sTKeXEqvMZOTR4"
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")


def fetch(url, timeout=60):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    return urllib.request.urlopen(req, timeout=timeout).read()


def probe(url):
    if not url:
        return "none"
    m = re.search(r"/spreadsheets/d/([A-Za-z0-9\-_]+)", url)
    if m:
        sid = m.group(1)
        g = re.search(r"[?&]gid=(\d+)", url)
        u = "https://docs.google.com/spreadsheets/d/%s/export?format=csv" % sid + ("&gid=%s" % g.group(1) if g else "")
        try:
            fetch(u, 40)
            return "sheet-public"
        except Exception:
            return "sheet-private"
    if "drive.google.com" in url:
        try:
            fetch(url, 40)
            return "drive-public"
        except Exception:
            return "drive-private"
    return "link"


def main():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(fetch("https://docs.google.com/spreadsheets/d/%s/export?format=xlsx" % SHEET)), data_only=True)
    out = {"checked_at": None, "sbus": []}
    from datetime import datetime
    out["checked_at"] = datetime.now().isoformat(timespec="seconds")
    for ws in wb.worksheets[1:]:
        name = ws.title
        if name.lower().startswith("00"):
            continue
        code = re.sub(r"^\d+\s*", "", name).strip()
        rows = []
        for r in range(10, ws.max_row + 1):
            cell = ws.cell(r, 6)
            target = cell.hyperlink.target if cell.hyperlink else None
            label = cell.value if isinstance(cell.value, str) else None
            ask = ws.cell(r, 1).value
            item = ws.cell(r, 3).value
            func = ws.cell(r, 2).value
            if not item and not target:
                continue
            url = target
            if not url and isinstance(label, str) and label.startswith("http"):
                url = label
            kind = "missing"
            if url:
                kind = "link"
            elif isinstance(label, str) and label.strip() and not label.strip().lower().startswith(("n/a", "erp ", "ibosddd")):
                kind = "named-file"
            rows.append({"ask": ask, "function": func, "item": str(item)[:120] if item else None,
                         "label": (label or "")[:120] or None, "url": url, "kind": kind})
        linked = [x for x in rows if x["kind"] in ("link", "named-file")]
        for x in linked:
            if x["url"]:
                x["access"] = probe(x["url"])
        out["sbus"].append({"sbu": code, "asks": len(rows), "linked": len(linked), "rows": linked})
    json.dump(out, open(os.path.join(DATA, "source_links.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    for s in out["sbus"]:
        pub = sum(1 for x in s["rows"] if x.get("access") in ("sheet-public", "drive-public"))
        print("%-24s asks=%3d linked=%2d public=%2d" % (s["sbu"], s["asks"], s["linked"], pub))
    return 0


if __name__ == "__main__":
    sys.exit(main())
