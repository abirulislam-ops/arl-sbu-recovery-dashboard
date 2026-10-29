#!/usr/bin/env python3
"""Watch the ARL marketing data tracker and report what changed.

  python scripts/watch_source.py            # check + log
  python scripts/watch_source.py --init     # create the baseline snapshot

The tracker is the Google Sheet "Tracker (Q1, FY 2026-27)". It is a live sheet:
SBUs in rows, data links in columns C-K. This script never invents data — it only
records additions, removals and re-pointed links, so the dashboard refresh can be
reviewed before anything reaches the MD.
"""
import hashlib, io, json, os, sys, urllib.request
from datetime import datetime

SHEET_ID = "1OXLfEnEdKVsTERVLZWtaaiUfTVaeTU_zhZcmEZCo6JQ"
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
SNAP = os.path.join(DATA, "_source_snapshot.json")
LOG = os.path.join(DATA, "_change_log.json")


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    return urllib.request.urlopen(req, timeout=60).read()


def read_tracker():
    import openpyxl
    blob = fetch("https://docs.google.com/spreadsheets/d/%s/export?format=xlsx" % SHEET_ID)
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    ws = wb[wb.sheetnames[0]]
    cols, links = {}, {}
    for c in range(3, 12):
        v = ws.cell(1, c).value
        if v:
            cols[openpyxl.utils.get_column_letter(c)] = str(v).strip()
    for r in range(2, ws.max_row + 1):
        sbu = ws.cell(r, 2).value
        if not sbu:
            continue
        for c in range(3, 12):
            cell = ws.cell(r, c)
            target = cell.hyperlink.target if cell.hyperlink else None
            label = cell.value if isinstance(cell.value, str) else None
            if target or (label and label not in ("Link", "")):
                links["%s|%s" % (str(sbu).strip(), openpyxl.utils.get_column_letter(c))] = {
                    "sbu": str(sbu).strip(),
                    "column": cols.get(openpyxl.utils.get_column_letter(c), ""),
                    "label": label,
                    "url": target,
                }
    return hashlib.sha256(blob).hexdigest(), cols, links


def main():
    init = "--init" in sys.argv
    digest, cols, links = read_tracker()
    now = datetime.now().isoformat(timespec="seconds")

    if init or not os.path.exists(SNAP):
        json.dump({"digest": digest, "checked_at": now, "columns": cols, "links": links},
                  open(SNAP, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        print("baseline written: %d links" % len(links))
        return 0

    prev = json.load(open(SNAP, encoding="utf-8"))
    old, new = prev.get("links", {}), links
    added = [k for k in new if k not in old]
    removed = [k for k in old if k not in new]
    changed = [k for k in new if k in old and (new[k].get("url") != old[k].get("url")
                                                or new[k].get("label") != old[k].get("label"))]
    if not (added or removed or changed) and digest == prev.get("digest"):
        print("no change (checked %s, %d links)" % (now, len(links)))
        return 0

    entry = {
        "checked_at": now,
        "sheet": "https://docs.google.com/spreadsheets/d/%s" % SHEET_ID,
        "added": [new[k] for k in added][:200],
        "removed": [old[k] for k in removed][:200],
        "changed": [{"sbu": new[k]["sbu"], "column": new[k]["column"],
                     "from": old[k].get("url") or old[k].get("label"),
                     "to": new[k].get("url") or new[k].get("label")} for k in changed][:200],
    }
    log = json.load(open(LOG, encoding="utf-8")) if os.path.exists(LOG) else []
    log.insert(0, entry)
    json.dump(log[:50], open(LOG, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    json.dump({"digest": digest, "checked_at": now, "columns": cols, "links": links},
              open(SNAP, "w", encoding="utf-8"), indent=1, ensure_ascii=False)

    print("CHANGE at %s — added %d, removed %d, changed %d" % (now, len(added), len(removed), len(changed)))
    for k in added[:12]:
        print("  + %-40s %s" % (k.replace("|", " · "), (new[k].get("label") or "")[:50]))
    for k in removed[:12]:
        print("  - %-40s %s" % (k.replace("|", " · "), (old[k].get("label") or "")[:50]))
    for k in changed[:12]:
        print("  ~ %-40s -> %s" % (k.replace("|", " · "), (new[k].get("url") or "")[:70]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
